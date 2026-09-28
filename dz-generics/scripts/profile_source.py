"""Profile the ministry nomenclature workbook for the normalization layer.

This is the inspection tool Task 2 couples to `normalize.py`. It reads the
main sheet (the one named ``Nomenclature *`` -- the other two sheets are
legacy lists whose header rows sit elsewhere, so they are out of scope), and
prints:

- the number of data rows,
- the descending frequency of ``form_key`` over those rows,
- the rows whose ``DOSAGE`` carries no number at all (``dose_key`` is None),
  with the raw values, and
- the *full* descending distribution of raw ``FORME`` strings, not a top-N.

The full raw FORME distribution is the authority for building
``FORM_CANON``: a truncated tail means unmapped forms. Re-run this script
whenever the ministry changes a column, and re-check the gates printed at the
bottom.

Run::

    python dz-generics/scripts/profile_source.py data/source/NOMENCLATURE.VERSION.AOUT_.2026-.xlsx
"""

from __future__ import annotations

import argparse
import sys
from collections import Counter

import openpyxl

from build_index import find_header  # column resolution (Task 1)
from normalize import dose_key, form_key


def _main_sheet(workbook: openpyxl.Workbook):
    """The sheet named ``Nomenclature *``, falling back to the first sheet."""
    for sheet in workbook.worksheets:
        if sheet.title.strip().lower().startswith("nomenclature"):
            return sheet
    return workbook.worksheets[0]


def _find_header(sheet):
    """Return (header_row_index, header_row) for the row containing ``FORME``.

    Scanning instead of hard-coding row 16 keeps the script correct when the
    ministry inserts or removes preamble rows.
    """
    for header_row, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if any(str(cell).strip().upper() == "FORME" for cell in row if cell is not None):
            return header_row, list(row)
    return None, None


def _data_rows(sheet):
    """Yield (row_number, row) for every non-empty row after the header row."""
    header_row, _ = _find_header(sheet)
    if header_row is None:
        raise SystemExit(f"no FORME header found in sheet {sheet.title!r}")

    for row_number, row in enumerate(
        sheet.iter_rows(min_row=header_row + 1, values_only=True),
        start=header_row + 1,
    ):
        if any(cell is not None for cell in row):
            yield row_number, row


def main(argv: list[str] | None = None) -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(
        description="Profile the ministry nomenclature workbook for normalize.py."
    )
    parser.add_argument("path", help="path to the ministry .xlsx file")
    args = parser.parse_args(argv)

    workbook = openpyxl.load_workbook(args.path, read_only=True)
    sheet = _main_sheet(workbook)

    rows = list(_data_rows(sheet))
    if not rows:
        raise SystemExit("no data rows found")

    header_row, headers = _find_header(sheet)
    form_col = find_header(headers, "FORME")
    dose_col = find_header(headers, "DOSAGE")
    if form_col is None or dose_col is None:
        raise SystemExit("FORME or DOSAGE column missing from the header row")

    print(f"sheet: {sheet.title}  (header row {header_row})")
    print(f"rows: {len(rows)}")
    print()

    forms = Counter()
    doses_missing: Counter[str] = Counter()
    for _, row in rows:
        form = str(row[form_col]).strip() if row[form_col] is not None else ""
        forms[form_key(form)] += 1
        dosage = row[dose_col]
        if dose_key(dosage) is None:
            doses_missing[str(dosage).strip() if dosage is not None else ""] += 1

    print("form_key (desc):")
    for key, count in sorted(forms.items(), key=lambda kv: (-kv[1], kv[0])):
        print(f"  {count:>5}  {key}")
    print()

    print(f"unparsed dosages: {sum(doses_missing.values())}")
    for raw, count in sorted(doses_missing.items(), key=lambda kv: (-kv[1], kv[0])):
        value = "''" if raw == "" else repr(raw)
        print(f"  {count:>5}  {value}")
    print()

    print("raw FORME distribution (desc, all values):")
    raw_forms = Counter()
    for _, row in rows:
        cell = row[form_col]
        raw_forms[str(cell).strip() if cell is not None else ""] += 1
    for raw, count in sorted(raw_forms.items(), key=lambda kv: (-kv[1], kv[0])):
        value = "''" if raw == "" else repr(raw)
        print(f"  {count:>5}  {value}")
    print()

    non_specifie = forms["NON_SPECIFIE"]
    share = non_specifie / len(rows) * 100
    print(
        f"NON_SPECIFIE: {non_specifie}/{len(rows)} = {share:.2f}%"
        "  (gate: < 5%)"
    )


if __name__ == "__main__":
    main()